"""Bounded local extraction, OCR, chunking and persistent FAISS document indexes."""

import asyncio
import hashlib
import os
import threading
import uuid
import zipfile
from datetime import UTC, datetime
from functools import lru_cache
from pathlib import Path

import numpy as np
from sqlalchemy import select
from starlette.concurrency import run_in_threadpool

from app.models.document_library import LibraryDocument as D

MODEL = "BAAI/bge-small-en-v1.5"
INDEX_ROOT = Path(os.getenv("DOCUMENT_INDEX_DIR", ".document-index")).resolve()
MODEL_CACHE_ROOT = INDEX_ROOT / "models"
MODEL_LOCK = threading.Lock()
MAX_TEXT = 2_000_000
MAX_PAGES = 300
INDEX_TIMEOUT_SECONDS = float(os.getenv("DOCUMENT_INDEX_TIMEOUT_SECONDS", "120"))


@lru_cache(maxsize=1)
def embedder():
    from fastembed import TextEmbedding

    return TextEmbedding(model_name=MODEL, cache_dir=str(MODEL_CACHE_ROOT), threads=2)


def embeddings(texts, query=False):
    with MODEL_LOCK:
        model = embedder()
        vectors = list(model.query_embed(texts) if query else model.embed(texts, batch_size=32))
    values = np.asarray(vectors, dtype="float32")
    import faiss

    faiss.normalize_L2(values)
    return values


@lru_cache(maxsize=1)
def ocr_engine():
    from rapidocr_onnxruntime import RapidOCR

    return RapidOCR(intra_op_num_threads=2, inter_op_num_threads=1)


def ocr(image):
    result, _ = ocr_engine()(image)
    return "\n".join(line[1] for line in (result or []))


def extract(path, filename):
    suffix = Path(filename).suffix.lower()
    sections = []
    limited = False

    def add(label, text):
        remaining = MAX_TEXT - sum(len(s["text"]) for s in sections)
        if remaining > 0 and text.strip():
            sections.append({"location": label, "text": text[:remaining]})

    if suffix in {".docx", ".xlsx", ".pptx", ".odt", ".ods", ".odp"}:
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100_000_000:
                raise ValueError("Document expands beyond the extraction limit")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("Password-protected PDF: upload an unlocked copy to index its text")
        limited = len(reader.pages) > MAX_PAGES
        for i, page in enumerate(reader.pages[:MAX_PAGES]):
            text = page.extract_text() or ""
            if len(text.strip()) < 20:
                import pypdfium2 as pdfium

                with pdfium.PdfDocument(path) as pdf:
                    pdf_page = pdf[i]
                    width, height = pdf_page.get_size()
                    bitmap = pdf_page.render(scale=min(1.5, 3000 / max(width, height)))
                    try:
                        text = ocr(np.asarray(bitmap.to_pil()))
                    finally:
                        bitmap.close()
            add(f"Page {i + 1}", text)
    elif suffix == ".docx":
        from docx import Document

        doc = Document(path)
        add("Document", "\n".join(p.text for p in doc.paragraphs))
        for i, table in enumerate(doc.tables):
            add(
                f"Table {i + 1}", "\n".join(" | ".join(c.text for c in r.cells) for r in table.rows)
            )
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book:
                lines = []
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i >= 20000:
                        break
                    lines.append(" | ".join(str(v) if v is not None else "" for v in row))
                add(sheet.title, "\n".join(lines))
        finally:
            book.close()
    elif suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(path)
        for sheet in book.sheets():
            add(
                sheet.name,
                "\n".join(
                    " | ".join(map(str, sheet.row_values(i)))
                    for i in range(min(sheet.nrows, 20000))
                ),
            )
    elif suffix == ".pptx":
        from pptx import Presentation

        for i, slide in enumerate(Presentation(path).slides):
            add(f"Slide {i + 1}", "\n".join(s.text for s in slide.shapes if s.has_text_frame))
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        from PIL import Image

        with Image.open(path) as image:
            image.thumbnail((3000, 3000))
            add("Image text", ocr(np.asarray(image.convert("RGB"))))
    elif suffix in {".odt", ".ods", ".odp"}:
        from bs4 import BeautifulSoup

        with zipfile.ZipFile(path) as archive:
            add("Document", BeautifulSoup(archive.read("content.xml"), "xml").get_text(" "))
    elif suffix == ".rtf":
        from striprtf.striprtf import rtf_to_text

        add("Document", rtf_to_text(path.read_text(errors="replace")))
    else:
        raw = path.read_bytes()[: MAX_TEXT * 2]
        if b"\x00" in raw[:8192]:
            return [], "This binary format is stored safely; text extraction is not supported."
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            return [], "Text encoding or file format is not supported."
        if suffix in {".html", ".htm", ".xml"}:
            from bs4 import BeautifulSoup

            text = BeautifulSoup(text, "html.parser").get_text(" ")
        add("Document", text)
    if limited or sum(len(section["text"]) for section in sections) >= MAX_TEXT:
        return (
            sections,
            "A large document was partially indexed (300 pages or 2 million characters).",
        )
    return sections, None if sections else "No readable text was found in this file."


def build_index(path, filename, document_id, organization_id):
    import faiss

    sections, warning = extract(path, filename)
    chunks = []
    for section in sections:
        words = section["text"].split()
        for start in range(0, len(words), 180):
            text = " ".join(words[start : start + 220])
            if text:
                chunks.append({"location": section["location"], "text": text})
    if not chunks:
        return dict(
            index_status="UNSUPPORTED",
            index_message=warning,
            extracted_text="",
            chunks=[],
            vector_path=None,
        )
    vectors = embeddings([c["text"] for c in chunks])
    index = faiss.IndexFlatIP(vectors.shape[1])
    index.add(vectors)
    folder = INDEX_ROOT / str(organization_id)
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"{document_id}-{uuid.uuid4().hex}.faiss"
    temporary = target.with_suffix(".tmp")
    faiss.write_index(index, str(temporary))
    os.replace(temporary, target)
    return dict(
        index_status="READY",
        index_message=warning,
        chunks=chunks,
        extracted_text="\n".join(s["text"] for s in sections),
        vector_path=str(target),
        content_hash=hashlib.sha256(path.read_bytes()).hexdigest(),
        indexed_at=datetime.now(UTC),
    )


@lru_cache(maxsize=128)
def load_vector_index(path):
    import faiss

    return faiss.read_index(path)


def search_vectors(documents, query):
    """Only authorized document indexes are loaded; snippets never cross an ACL boundary."""

    vector = embeddings([query], query=True)
    hits = {}
    for document in documents:
        if not document.vector_path or document.index_status != "READY":
            continue
        path = Path(document.vector_path).resolve()
        if not path.is_relative_to(INDEX_ROOT) or not path.is_file():
            continue
        index = load_vector_index(str(path))
        scores, ids = index.search(vector, 1)
        position = int(ids[0][0])
        if 0 <= position < len(document.chunks):
            hits[document.id] = (float(scores[0][0]), document.chunks[position])
    return hits


async def index_one(session_factory, storage):
    doc_info = None
    async with session_factory() as session:
        row = await session.scalar(
            select(D)
            .where(D.index_status == "PENDING", D.is_active.is_(True))
            .order_by(D.created_at)
            .with_for_update(skip_locked=True)
            .limit(1)
        )
        if row is None:
            return False
        row.index_status = "PROCESSING"
        await session.commit()
        doc_info = (row.id, row.storage_path, row.file_name, row.organization_id)

    doc_id, storage_path, file_name, org_id = doc_info
    try:
        path = await run_in_threadpool(storage.resolve, storage_path)
        result = await asyncio.wait_for(
            run_in_threadpool(build_index, path, file_name, doc_id, org_id),
            timeout=INDEX_TIMEOUT_SECONDS,
        )
    except Exception as error:
        from pypdf.errors import PdfReadError

        message = (
            "Text indexing failed. Retry after checking the file and local model availability."
        )
        if isinstance(error, TimeoutError | asyncio.TimeoutError):
            message = f"Text indexing timed out after {INDEX_TIMEOUT_SECONDS:g} seconds. Check file size or model availability."
        elif isinstance(error, PdfReadError):
            message = (
                "This PDF is damaged or incomplete. Upload a valid PDF to extract its text."
            )
        elif isinstance(error, ValueError):
            message = str(error)[:400]
        result = {
            "index_status": "FAILED",
            "index_message": message,
        }
        import structlog

        structlog.get_logger().warning(
            "document_index_failed",
            document_id=str(doc_id),
            file_name=file_name,
            error_type=type(error).__name__,
        )

    async with session_factory() as session:
        row = await session.get(D, doc_id)
        if row:
            for key, value in result.items():
                setattr(row, key, value)
            await session.commit()
    return True


async def worker(app):
    import structlog

    while True:
        try:
            if not await index_one(app.state.session_factory, app.state.storage):
                await asyncio.sleep(5)
            else:
                await asyncio.sleep(0.5)
        except asyncio.CancelledError:
            raise
        except Exception:
            structlog.get_logger().exception("document_worker_failed")
            await asyncio.sleep(15)
